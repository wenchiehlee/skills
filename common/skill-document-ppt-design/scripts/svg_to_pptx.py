#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
svg_to_pptx.py — 將一組 SVG 頁面批次轉換為 PowerPoint (.pptx)

流程：
  1. 用 Playwright + Chromium headless 將每頁 SVG 截圖為 PNG
  2. 用 python-pptx 將 PNG 組合為 16:9 (1280×720) 的 PPTX

使用方式：
  python scripts/svg_to_pptx.py --slides slides/ --output deck.pptx
  python scripts/svg_to_pptx.py --slides slides/ --output deck.pptx --width 1920 --height 1080

依賴：
  pip install playwright python-pptx pillow
  playwright install chromium
"""
from __future__ import annotations

import argparse
import sys
import tempfile
from pathlib import Path

# ---------------------------------------------------------------------------
# 依賴檢查
# ---------------------------------------------------------------------------

def _check_deps() -> None:
    missing = []
    try:
        import playwright  # noqa: F401
    except ImportError:
        missing.append("playwright (pip install playwright && playwright install chromium)")
    try:
        import pptx  # noqa: F401
    except ImportError:
        missing.append("python-pptx (pip install python-pptx)")
    try:
        import PIL  # noqa: F401
    except ImportError:
        missing.append("pillow (pip install pillow)")
    if missing:
        print("[Error] 缺少依賴套件：")
        for m in missing:
            print(f"  - {m}")
        sys.exit(1)


# ---------------------------------------------------------------------------
# SVG → PNG
# ---------------------------------------------------------------------------

def svg_to_png(svg_path: Path, png_path: Path, width: int, height: int) -> None:
    """使用 Playwright 將 SVG 截圖為 PNG。"""
    from playwright.sync_api import sync_playwright

    svg_uri = svg_path.resolve().as_uri()
    with sync_playwright() as p:
        browser = p.chromium.launch(headless=True)
        page = browser.new_page(viewport={"width": width, "height": height})
        page.goto(svg_uri, wait_until="networkidle")
        page.screenshot(path=str(png_path), full_page=False)
        browser.close()


# ---------------------------------------------------------------------------
# PNG → PPTX
# ---------------------------------------------------------------------------

def pngs_to_pptx(
    png_paths: list[Path],
    output_path: Path,
    slide_width_px: int,
    slide_height_px: int,
) -> None:
    """將 PNG 清單組合為 PPTX（每頁一張圖片佔滿投影片）。"""
    from pptx import Presentation
    from pptx.util import Emu

    # PowerPoint 內部使用 EMU（English Metric Units）
    # 1 inch = 914400 EMU；1280×720 px @96 dpi ≈ 13.33" × 7.5"
    SLIDE_W = Emu(int(slide_width_px / 96 * 914400))
    SLIDE_H = Emu(int(slide_height_px / 96 * 914400))

    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    blank_layout = prs.slide_layouts[6]  # 空白版面

    for png in png_paths:
        slide = prs.slides.add_slide(blank_layout)
        slide.shapes.add_picture(str(png), 0, 0, SLIDE_W, SLIDE_H)

    prs.save(str(output_path))
    print(f"[OK] PPTX saved → {output_path}")


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

def main() -> None:
    _check_deps()

    parser = argparse.ArgumentParser(
        description="SVG slides → PPTX"
    )
    parser.add_argument("--slides", required=True, help="SVG 檔案所在目錄")
    parser.add_argument("--output", required=True, help="輸出的 .pptx 路徑")
    parser.add_argument("--width", type=int, default=1280, help="截圖寬度（default: 1280）")
    parser.add_argument("--height", type=int, default=720, help="截圖高度（default: 720）")
    args = parser.parse_args()

    slides_dir = Path(args.slides)
    if not slides_dir.is_dir():
        print(f"[Error] --slides 目錄不存在：{slides_dir}")
        sys.exit(1)

    svg_files = sorted(slides_dir.glob("*.svg"))
    if not svg_files:
        print(f"[Error] 在 {slides_dir} 找不到任何 .svg 檔案")
        sys.exit(1)

    print(f"[Info] 找到 {len(svg_files)} 個 SVG 頁面")

    output_path = Path(args.output)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    with tempfile.TemporaryDirectory() as tmpdir:
        tmp = Path(tmpdir)
        png_paths: list[Path] = []

        for i, svg in enumerate(svg_files, start=1):
            png = tmp / f"{i:03d}.png"
            print(f"  [{i:02d}/{len(svg_files)}] {svg.name} → {png.name}")
            svg_to_png(svg, png, args.width, args.height)
            png_paths.append(png)

        pngs_to_pptx(png_paths, output_path, args.width, args.height)


if __name__ == "__main__":
    main()
